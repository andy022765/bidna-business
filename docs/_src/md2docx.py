#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown → .docx для наших рабочих документов. Заголовки, таблицы, цитаты, списки."""
import re, sys, io
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
ACCENT = RGBColor(0x2B, 0x3A, 0x7A)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def left_bar(par, color="2B3A7A"):
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single'); left.set(qn('w:sz'), '18')
    left.set(qn('w:space'), '10'); left.set(qn('w:color'), color)
    bdr.append(left); pPr.append(bdr)

TOK = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|~~[^~]+~~)')

def runs(par, text, base_size=11, color=INK, bold_all=False):
    for part in TOK.split(text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1]); r.font.name = 'Menlo'; r.font.size = Pt(base_size-1.5)
            r.font.color.rgb = RGBColor(0x30,0x30,0x30)
        elif part.startswith('~~') and part.endswith('~~'):
            r = par.add_run(part[2:-2]); r.font.strike = True; r.font.color.rgb = MUTED
        elif part.startswith('*') and part.endswith('*'):
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            r = par.add_run(part)
        if r.font.size is None: r.font.size = Pt(base_size)
        if r.font.color.rgb is None: r.font.color.rgb = color
        if bold_all: r.bold = True

def novaya_numeraciya(doc, start):
    """Свой счётчик для нового нумерованного списка, с номера из markdown. Без этого все пункты
    'List Number' в документе идут одним счётом, и «п. 6» в тексте не совпадает с номером в Word."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    style_num = doc.styles['List Number'].element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(style_num).abstractNumId.val
    num = numbering.add_num(abstract)
    num.add_lvlOverride(ilvl=0).add_startOverride(start)
    return num.numId

def build(md_path, out_path):
    lines = io.open(md_path, encoding='utf-8').read().split('\n')
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    for s in ('Heading 1','Heading 2','Heading 3'):
        f = doc.styles[s].font; f.name = 'Calibri'; f.color.rgb = ACCENT
    doc.styles['Heading 1'].font.size = Pt(20)
    doc.styles['Heading 2'].font.size = Pt(15)
    doc.styles['Heading 3'].font.size = Pt(12.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)

    spisok = {}                      # текущий нумерованный список: его numId и последний пункт
    i, n = 0, len(lines)
    while i < n:
        ln = lines[i]
        s = ln.strip()

        if not s:
            i += 1; continue

        # горизонтальная линия
        if re.fullmatch(r'-{3,}', s):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(2)
            pPr = p._p.get_or_add_pPr(); bdr = OxmlElement('w:pBdr')
            bot = OxmlElement('w:bottom')
            bot.set(qn('w:val'),'single'); bot.set(qn('w:sz'),'6')
            bot.set(qn('w:space'),'1'); bot.set(qn('w:color'),'CCCCCC')
            bdr.append(bot); pPr.append(bdr)
            i += 1; continue

        # заголовки
        m = re.match(r'^(#{1,4})\s+(.*)$', s)
        if m:
            lvl = min(len(m.group(1)), 3)
            p = doc.add_paragraph(style=f'Heading {lvl}')
            txt = re.sub(r'\*\*(.+?)\*\*', r'\1', m.group(2))
            r = p.add_run(txt)
            p.paragraph_format.space_before = Pt(16 if lvl <= 2 else 10)
            p.paragraph_format.space_after = Pt(5)
            i += 1; continue

        # таблица
        if s.startswith('|') and i+1 < n and re.match(r'^\|[\s:|-]+\|$', lines[i+1].strip()):
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                row = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not re.match(r'^[\s:|-]+$', ''.join(row)): rows.append(row)
                i += 1
            cols = max(len(r) for r in rows)
            t = doc.add_table(rows=0, cols=cols); t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            for ri, row in enumerate(rows):
                cells = t.add_row().cells
                for ci in range(cols):
                    cell = cells[ci]
                    cell.paragraphs[0].text = ''
                    val = row[ci] if ci < len(row) else ''
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_after = Pt(2)
                    par.paragraph_format.space_before = Pt(2)
                    runs(par, val, base_size=9.5, bold_all=(ri == 0))
                    if ri == 0: shade(cell, 'EDEFF5')
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # цитата — блок сообщений
        if s.startswith('>'):
            block = []
            while i < n and lines[i].strip().startswith('>'):
                block.append(re.sub(r'^\s*>\s?', '', lines[i]))
                i += 1
            for b in block:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.7)
                p.paragraph_format.space_after = Pt(3)
                left_bar(p)
                if b.strip():
                    runs(p, b.strip(), base_size=10.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # списки
        m = re.match(r'^(\s*)[-*]\s+(.*)$', ln)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            if len(m.group(1)) >= 2: p.paragraph_format.left_indent = Cm(1.6)
            runs(p, m.group(2)); i += 1; continue
        m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', ln)
        if m:
            pred = doc.paragraphs[-1]._p if doc.paragraphs else None
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            if pred is None or pred is not spisok.get('posledniy'):   # перед пунктом не пункт — новый список
                spisok['numId'] = novaya_numeraciya(doc, int(m.group(2)))
            numPr = p._p.get_or_add_pPr().get_or_add_numPr()
            numPr.get_or_add_ilvl().val = 0
            numPr.get_or_add_numId().val = spisok['numId']
            spisok['posledniy'] = p._p
            runs(p, m.group(3)); i += 1; continue

        # обычный абзац — склеиваем до пустой строки
        buf = []
        while i < n and lines[i].strip() and not re.match(r'^(#{1,4}\s|\||>|\s*[-*]\s|\s*\d+\.\s|-{3,}$)', lines[i].strip()):
            buf.append(lines[i].strip()); i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        runs(p, ' '.join(buf))

    doc.save(out_path)
    return out_path

if __name__ == '__main__':
    print(build(sys.argv[1], sys.argv[2]))
